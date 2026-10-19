"""Glue between the Streamlit UI and the ceda-column-design engine (engine/ = skill scripts, not re-implemented).

Pipeline:  uploads -> ETABS tables per model (merged) -> normalised workbook per model (TABLE: blocks)
           -> config.json -> engine run.py -> engine make_outputs.py (--per-mark) -> results for the UI.
"""
import os, sys, json, re, io, math, zipfile, pickle, statistics, contextlib
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE = os.path.join(HERE, 'engine')
if ENGINE not in sys.path:
    sys.path.insert(0, ENGINE)
import etabs_io as E          # noqa: E402
import layout as LY           # noqa: E402
import design as DS           # noqa: E402

ENGINE_VERSION = open(os.path.join(ENGINE, 'ENGINE_VERSION')).read().strip()
LAP_KNOWN = {400: {25: 1550, 28: 1875, 32: 2325}, 560: {25: 1325, 28: 1600, 32: 1975}}   # CEDA S-00-04-04 (confirmed values)
NEEDED = {
    'Element Forces - Columns': 'forces (required)',
    'Load Combination Definitions': 'combo factors – ข้อ 7 check, seismic/gravity split, basement factors',
    'Column Object Connectivity': 'Lu per column-story, story order',
    'Point Object Connectivity': 'XY (label matching between models), elevations',
    'Frame Assignments - Section Properties': 'section of every column-story',
    'Frame Section Property Definitions - Concrete Rectangular': 'section dimensions',
    'Story Forces': 'sway check Q (seismic model)',
    'Story Drifts': 'sway check Q (seismic model)',
}


# ------------------------------------------------------------------------------------------------ reading uploads
def _headerless_combos(path):
    """Load Combination Definitions pasted without the TABLE:/header rows (Name, Type, Auto, Load Name, Mode, SF)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows, cur = [], None
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            r = list(r) + [None] * 6
            if r[0] is not None and str(r[0]).upper().startswith('TABLE:'):
                return []
            ld, sf = r[3], r[5]
            if isinstance(ld, str) and isinstance(sf, (int, float)):
                cur = r[0] if r[0] not in (None, '') else cur
                if cur:
                    rows.append({'Name': str(cur), 'Load Name': ld.strip(), 'SF': float(sf)})
    return rows if len(rows) >= 2 else []


def read_upload(path):
    """-> (tables, units, info). Accepts normal ETABS exports and headerless combo definitions."""
    tables, units = E.read_tables(path)
    info = []
    if not tables:
        rows = _headerless_combos(path)
        if rows:
            tables = {'Load Combination Definitions': rows}
            units = {'Load Combination Definitions': {}}
            info.append(f'{os.path.basename(path)}: read as Load Combination Definitions without header ({len(rows)} rows)')
        else:
            info.append(f'{os.path.basename(path)}: no ETABS "TABLE:" block found – ignored')
    return tables, units, info


def merge(parts):
    T, U = {}, {}
    for t, u in parts:
        for k, v in t.items():
            T.setdefault(k, []).extend(v)
            U.setdefault(k, u.get(k, {}))
    return T, U


def write_tables(path, tables, units):
    """write tables back as an ETABS-style workbook (one sheet, TABLE: blocks) for the engine reader."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for name, rows in tables.items():
        if not rows:
            continue
        hdr = []
        for r in rows:
            for k in r:
                if k not in hdr:
                    hdr.append(k)
        ws.append(['TABLE:  ' + name])
        ws.append(hdr)
        u = units.get(name, {})
        ws.append([u.get(h, '') for h in hdr])
        for r in rows:
            ws.append([r.get(h) for h in hdr])
    wb.save(path)


def inventory(tables):
    out = []
    for name, why in NEEDED.items():
        t, real = E.find_table(tables, name)
        out.append(dict(table=name, found=bool(t), rows=len(t) if t else 0, use=why))
    return out


def scan(tables, units, gravity=None):
    """labels, stories, combos and what is already known (sections, Lu, elevations) in one model."""
    gn = gravity or E.GRAV_DEFAULT
    allg = {E._norm(x): g for g, xs in gn.items() for x in xs}
    t, n = E.find_table(tables, 'Element Forces - Columns')
    forces = E.column_forces(tables, units, 'x') if t else []
    labels = sorted({f['label'] for f in forces}, key=lambda s: (len(s), s))
    stories = sorted({f['story'] for f in forces})
    combos = sorted({f['combo'] for f in forces if 'comb' in f['ctype'].lower() or not f['ctype'].strip()})
    cases = sorted({f['combo'] for f in forces} - set(combos))
    fac, kind = E.combo_factors(tables, gn)
    case_names = sorted({c for f in fac.values() for c in f}) or cases
    geo = E.geometry(tables, units)
    secs = E.sections(tables, units)
    order = E.story_order(geo) if geo else []
    from_z = bool(order)
    if not order:
        from run import story_rank
        order = sorted(stories, key=story_rank)
    order = [s for s in order if s in stories] + [s for s in stories if s not in order]
    L, z = {}, {}
    for (lab, st), g in geo.items():
        if 'L' in g:
            L.setdefault(st, []).append(g['L'] / 100)
        if 'zb' in g:
            z[st] = (round(g['zb'] / 100, 3), round(g['zt'] / 100, 3))
    Lu = {st: round(statistics.median(v), 3) for st, v in L.items()}
    sec_of = {}
    for (lab, st), g in geo.items():
        s = g.get('section')
        d = (secs.get(s) or E.section_from_name(s)) if s else None
        if d:
            sec_of.setdefault(lab, set()).add((d['shape'], round(d['t2'] * 10), round(d['t3'] * 10), bool(d.get('assumed'))))
    has_cases = any(allg.get(E._norm(f['combo'])) for f in forces)
    return dict(labels=labels, stories=order, combos=combos, cases=cases, case_names=case_names, defined=sorted(fac), kind=kind,
                Lu=Lu, z=z, sec_of=sec_of, sec_labels=set(sec_of), L_keys={k for k, g in geo.items() if 'L' in g},
                order_from_z=from_z, has_gravity_cases=has_cases, n_forces=len(forces))


def synth_combos(rows):
    """designer-entered factors -> Load Combination Definitions rows (case names the engine recognises)."""
    out = []
    for r in rows:
        name = str(r['combo'])
        for case, key in (('DL', 'DL'), ('SDL', 'SDL'), ('LL', 'LL'), ('EQ', 'E'), ('WIND', 'W')):
            v = r.get(key)
            if v not in (None, '') and not (isinstance(v, float) and math.isnan(v)) and float(v) != 0:
                out.append({'Name': name, 'Load Name': case, 'SF': float(v)})
    return out


# ------------------------------------------------------------------------------------------------ run
def run_engine(workdir, cfg, log=None):
    """engine run.py + make_outputs.py (per-mark files). Returns (results dict, outdir)."""
    import importlib
    import run as RUN
    import make_outputs as MO
    importlib.reload(RUN); importlib.reload(MO)
    cp = os.path.join(workdir, 'config.json')
    json.dump(cfg, open(cp, 'w'), ensure_ascii=False, indent=1)
    out = os.path.join(workdir, 'OUT')
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        RUN.main(cp, out)
        MO.main(out, per_mark=True)
    if log is not None:
        log.append(buf.getvalue())
    return json.load(open(os.path.join(out, 'results.json'))), out


def detail_excel(outdir, mark, bottom_story):
    import make_outputs as MO
    with contextlib.redirect_stdout(io.StringIO()):
        MO.main(outdir, (mark, bottom_story))
    job = json.load(open(os.path.join(outdir, 'results.json')))['config']['project'].get('job', 'job')
    return os.path.join(outdir, f'{job}_{mark}_{bottom_story}_Detailed_Calc.xlsx')


def files(outdir):
    res = json.load(open(os.path.join(outdir, 'results.json')))
    job = res['config']['project'].get('job', 'job')
    main = [(f'{job}_Column_Decision.xlsx', 'Decision table (label × story, f\'c study, notes, schedule)'),
            (f'{job}_Column_Calc.pdf', 'Calculation sheets – all marks'),
            (f'{job}_Column_Backup.xlsx', 'Back-up calculation – all marks (live formulas)'),
            (f'{job}_Table_of_Column_DRAFT.dxf', 'Table of Column – DXF (draft)'),
            (f'{job}_Table_of_Column_DRAFT.pdf', 'Table of Column – PDF (draft)'),
            (f'{job}_Column_Schedule_DRAFT.xlsx', 'Schedule with bar coordinates (Revit)'),
            ('results.json', 'Raw engine results')]
    main = [(os.path.join(outdir, f), f, d) for f, d in main if os.path.exists(os.path.join(outdir, f))]
    md = os.path.join(outdir, 'marks')
    per = {}
    if os.path.isdir(md):
        for f in sorted(os.listdir(md)):
            m = re.match(rf'{re.escape(job)}_(.+?)_(Calc\.pdf|Backup\.xlsx)$', f)
            if m:
                per.setdefault(m.group(1), {})[m.group(2)] = os.path.join(md, f)
    return main, per


def zip_all(outdir):
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, 'w', zipfile.ZIP_DEFLATED) as z:
        for root, _, fs in os.walk(outdir):
            for f in fs:
                if f.endswith('.pkl'):
                    continue
                p = os.path.join(root, f)
                z.write(p, os.path.relpath(p, outdir))
    return bio.getvalue()


# ------------------------------------------------------------------------------------------------ views
def load_state(outdir):
    return pickle.load(open(os.path.join(outdir, 'state.pkl'), 'rb'))


def mark_band_eval(outdir, mark, band):
    """engine evaluation of every label-story of one mark-band (same functions as the calc sheets)."""
    res = json.load(open(os.path.join(outdir, 'results.json')))
    st = load_state(outdir)
    cache = DS.Cache()
    fy = st['cfg']['materials'].get('fy', 5000)
    items = []
    for r in res['results']:
        if r['mark'] != mark or tuple(r['band']) != tuple(band):
            continue
        for s in r['stories']:
            c = st['cols'][(r['label'], s)]
            ev = DS.evaluate(cache, st['grids'][c['gkey']], c, st['fc_of'][s], fy, r['db'], tuple(r['counts']))
            items.append(dict(label=r['label'], story=s, r=r, c=c, ev=ev))
    return items


def surface_figures(items):
    """plotly: 3D capped φ-surface + demand points, P–Mx, P–My, Mx–My contour at governing Pu (units t, t-m)."""
    import numpy as np
    import plotly.graph_objects as go
    gov = max(items, key=lambda q: q['ev']['dcmax'])
    S = gov['ev']['S']
    V = S.hull.points * S._scale()
    I, J, K = S.hull.simplices.T
    f3 = go.Figure()
    f3.add_trace(go.Mesh3d(x=V[:, 1] / 1e3, y=V[:, 2] / 1e3, z=V[:, 0] / 1e3, i=I, j=J, k=K, opacity=0.35, color='#8fb3d9',
                           name='φ-surface', hoverinfo='skip'))
    for q in items:
        mg, P = q['ev']['mg'], q['c']['rows']['P']
        dc = q['ev']['dc']
        txt = [f"{q['label']} {q['story']}<br>{m['combo']}<br>D/C {d:.3f}" for m, d in zip(q['c']['meta'], dc)]
        f3.add_trace(go.Scatter3d(x=mg['x']['Md'] / 1e3, y=mg['y']['Md'] / 1e3, z=P / 1e3, mode='markers',
                                  marker=dict(size=3, color=dc, colorscale='RdYlGn_r', cmin=0, cmax=1.1),
                                  text=txt, hoverinfo='text', name=f"{q['label']} {q['story']}"))
    i = gov['ev']['gov']
    gm = gov['ev']['mg']
    gp = (gm['x']['Md'][i] / 1e3, gm['y']['Md'][i] / 1e3, gov['c']['rows']['P'][i] / 1e3)
    f3.add_trace(go.Scatter3d(x=[0, gp[0]], y=[0, gp[1]], z=[0, gp[2]], mode='lines+markers', line=dict(color='red', width=5),
                              name=f"governing ray D/C {gov['ev']['dcmax']:.3f}"))
    f3.update_layout(scene=dict(xaxis_title='Mx (t-m)', yaxis_title='My (t-m)', zaxis_title='P (t)'), height=620,
                     margin=dict(l=0, r=0, t=30, b=0), legend=dict(font=dict(size=9)))
    figs2 = []
    for ax, key, lab in (('x', 'M3', 'Mx'), ('y', 'M2', 'My')):
        C = S.uniaxial(ax)
        f = go.Figure()
        f.add_trace(go.Scatter(x=C[:, 1] / 1e3, y=C[:, 0] / 1e3, mode='lines', name='φPn–φMn'))
        f.add_trace(go.Scatter(x=-C[:, 1] / 1e3, y=C[:, 0] / 1e3, mode='lines', showlegend=False, line=dict(color='#636efa')))
        for q in items:
            f.add_trace(go.Scatter(x=q['ev']['mg'][ax]['Md'] / 1e3, y=q['c']['rows']['P'] / 1e3, mode='markers', marker=dict(size=5),
                                   name=f"{q['label']} {q['story']}"))
        f.update_layout(xaxis_title=f'{lab} magnified (t-m)', yaxis_title='P (t)', height=420, margin=dict(l=10, r=10, t=30, b=10),
                        title=f'P–{lab} (uniaxial section of the surface)')
        figs2.append(f)
    Pg = min(gov['c']['rows']['P'][i], S.phiPnmax * 0.9999)
    Cc = S.contour(Pg)
    fc = go.Figure()
    if len(Cc):
        fc.add_trace(go.Scatter(x=Cc[:, 0] / 1e3, y=Cc[:, 1] / 1e3, mode='lines', name=f'contour at Pu = {Pg/1e3:.0f} t'))
    fc.add_trace(go.Scatter(x=[gp[0]], y=[gp[1]], mode='markers', marker=dict(size=10, color='red'), name='governing demand'))
    fc.update_layout(xaxis_title='Mx (t-m)', yaxis_title='My (t-m)', height=420, margin=dict(l=10, r=10, t=30, b=10),
                     title='Mx–My at governing Pu', yaxis_scaleanchor='x')
    return f3, figs2, fc, gov


def dxf_png(path, dpi=200):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import ezdxf
    from ezdxf.addons.drawing import RenderContext, Frontend
    from ezdxf.addons.drawing.matplotlib import MatplotlibBackend
    doc = ezdxf.readfile(path)
    msp = doc.modelspace()
    from ezdxf import bbox
    ext = bbox.extents(msp)
    w, h = (ext.size.x, ext.size.y) if ext.has_data else (16, 11)
    fig = plt.figure(figsize=(24 * w / max(w, h) + 0.5, 24 * h / max(w, h) + 0.5))
    ax = fig.add_axes([0, 0, 1, 1])
    from ezdxf.addons.drawing.config import Configuration, BackgroundPolicy, ColorPolicy
    cfg = Configuration(background_policy=BackgroundPolicy.WHITE, color_policy=ColorPolicy.BLACK)
    Frontend(RenderContext(doc), MatplotlibBackend(ax), config=cfg).draw_layout(msp, finalize=True)
    bio = io.BytesIO()
    fig.savefig(bio, format='png', dpi=dpi, facecolor='white')
    plt.close(fig)
    return bio.getvalue()


# ------------------------------------------------------------------------------------------------ app pipeline (shared by UI and tests)
def prepare(up, tpl):
    """up: {'seismic'|'full_reduced'|'full_unreduced': (tables, units) or None}; tpl: inputs.read_template() dict."""
    import inputs as INP
    single = up.get('full_reduced') is None or up.get('seismic') is None
    base = {'seismic': up.get('seismic') or up.get('full_reduced'), 'full_reduced': up.get('full_reduced') or up.get('seismic'),
            'full_unreduced': up.get('full_unreduced')}
    label_story = set()
    for m in base.values():
        if m:
            label_story |= {(f['label'], f['story']) for f in E.column_forces(m[0], m[1], 'x')}
    eff, rep = {}, []
    for k, m in base.items():
        if m is None:
            eff[k] = None; continue
        T, U, r = INP.apply(m[0], m[1], tpl, label_story)
        eff[k] = (T, U); rep += [f'{k}: {x}' for x in r]
    cfg = INP.to_cfg(tpl)
    gn = cfg['gravity_cases']
    scS, scR = scan(*eff['seismic'], gravity=gn), scan(*eff['full_reduced'], gravity=gn)
    labels = sorted({l for l, s in label_story}, key=lambda s: (len(s), s))
    stories = cfg['story_order'] or (scR['stories'] + [s for s in scS['stories'] if s not in scR['stories']])
    combos = sorted(set(scS['combos']) | set(scR['combos']))
    undefined = [c for c in combos if c not in set(scS['defined']) | set(scR['defined'])]
    found = dict(labels=labels, stories=stories, combos=combos, undefined=undefined or combos, undefined_only=undefined,
                 cases=scR['case_names'] or scS['case_names'], sec_of={**scS['sec_of'], **scR['sec_of']}, z={**scS['z'], **scR['z']})
    return dict(single=single, eff=eff, merge_rep=rep, cfg_tpl=cfg, scS=scS, scR=scR, labels=labels, stories=stories,
                combos=combos, undefined=undefined, found=found, label_story=label_story)


def build_cfg(prep, workdir):
    import inputs as INP
    paths = {}
    for k, fn in (('seismic', 'seismic.xlsx'), ('full_reduced', 'full_reduced.xlsx'), ('full_unreduced', 'full_unreduced.xlsx')):
        if prep['eff'][k]:
            write_tables(os.path.join(workdir, fn), *prep['eff'][k]); paths[k] = fn
    cfg = json.loads(json.dumps(prep['cfg_tpl']))
    if not cfg['lap_table']:
        cfg['lap_table'] = {str(f): {str(d): v for d, v in t.items()} for f, t in INP.LAP_KNOWN.items()}
    cfg['inputs'] = paths
    if not prep['single']:
        cfg['basement_stories'] = []
    if not cfg['story_order']:
        cfg['story_order'] = prep['stories']
    return cfg
