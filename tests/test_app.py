"""App tests – run:  python3 tests/test_app.py   -> ALL PASS
A  engine self-test (skill scripts bundled in engine/) must print ALL PASS
B  parity: app pipeline (upload reader -> normalised workbook -> engine) == engine run directly on the ETABS files
C  one-model upload + combo sheet without header (VOCO test files format) -> runs, combo check flags wrong factors,
   every output file + per-mark calc/back-up exist, DXF renders, interaction surface builds
"""
import os, sys, json, subprocess, tempfile, shutil, copy
HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)
import app_io as AIO  # noqa: E402

fails = []


def check(name, cond, info=''):
    print(('PASS ' if cond else 'FAIL ') + name + (f'  [{info}]' if info else ''))
    if not cond:
        fails.append(name)


ENG = os.path.join(HERE, 'engine')
tmp = tempfile.mkdtemp()
# ---------------------------------------------------------------- A
r = subprocess.run([sys.executable, '-I', os.path.join(ENG, 'selftest.py'), ENG], cwd=tmp, capture_output=True, text=True)
check('A engine self-test ALL PASS', 'ALL PASS' in r.stdout, r.stdout.strip().splitlines()[-1] if r.stdout else r.stderr[-300:])

# ---------------------------------------------------------------- B parity on the synthetic 3-model ETABS set
subprocess.run([sys.executable, '-I', os.path.join(ENG, 'selftest_data', 'make_synthetic_etabs.py'), tmp], check=True)
base = json.load(open(os.path.join(ENG, 'config_template.json')))
base['project'] = dict(name='PARITY', job='PAR')
base['seismic']['frame'] = 'gravity'
base['slender']['overrides'] = []
base['exclude_combos'] = ['U_B_99']
cfgA = copy.deepcopy(base)
cfgA['inputs'] = dict(seismic='seis.xlsx', full_reduced='full_red.xlsx', full_unreduced='full_unred.xlsx')
dA = os.path.join(tmp, 'A'); os.makedirs(dA)
for f in ('seis.xlsx', 'full_red.xlsx', 'full_unred.xlsx'):
    shutil.copy(os.path.join(tmp, f), dA)
resA, outA = AIO.run_engine(dA, cfgA)
dB = os.path.join(tmp, 'B'); os.makedirs(dB)
paths = {}
for key, f in (('seismic', 'seis.xlsx'), ('full_reduced', 'full_red.xlsx'), ('full_unreduced', 'full_unred.xlsx')):
    t, u, _ = AIO.read_upload(os.path.join(tmp, f))
    AIO.write_tables(os.path.join(dB, key + '.xlsx'), t, u)
    paths[key] = key + '.xlsx'
cfgB = copy.deepcopy(base); cfgB['inputs'] = paths
resB, outB = AIO.run_engine(dB, cfgB)
sig = lambda R: sorted((r['label'], tuple(r['band']), r['mark'], r['nbar'], r['db'], round(max(q['dc'] for q in r['per']), 4)) for r in R['results'])
check('B app pipeline == engine on the original files (labels, ranges, marks, bars, D/C)', sig(resA) == sig(resB), str(sig(resB))[:200])
check('B basement assembled through the app path', bool(resB['input_report'].get('basement_stories')), str(resB['input_report'].get('basement_stories')))

# ---------------------------------------------------------------- C one model, everything else from the CEDA input workbook
import inputs as INP
import openpyxl
tdir = sys.argv[1] if len(sys.argv) > 1 else None
if tdir and os.path.exists(os.path.join(tdir, 'Column Load.xlsx')):
    t, u, info = AIO.read_upload(os.path.join(tdir, 'Column Load.xlsx'))
    empty = INP.read_template(os.path.join(tmp, 'none.xlsx')) if False else dict(criteria={}, stories=[], columns=[], sections=[],
              combos_rows=[], combo_cases=[], lap_rows=[], exclude=[], approve=[], issues=[])
    p0 = AIO.prepare({'seismic': (t, u)}, empty)
    check('C0 forces only: combos, sizes and Lu reported missing', bool(p0['undefined']) and not p0['scS']['sec_labels'] and not p0['scS']['L_keys'])
    # prefilled workbook -> fill it like the designer would in Excel
    tp = os.path.join(tmp, 'CEDA_Column_Input.xlsx')
    open(tp, 'wb').write(INP.make_template(p0['found']))
    wb = openpyxl.load_workbook(tp)
    cr = {r[0].value: r[1] for r in wb['Criteria'].iter_rows(min_row=3)}
    for k, v in {'project.job': 'VT', 'project.name': 'VOCO test', 'slender.beta_d': 0.888, 'slender.k': 0.9, 'materials.fc_mode': 'fixed',
                 'materials.fc_bottom': 560, 'materials.fc_reduce_at': '4th'}.items():
        cr[k].value = v
    ws = wb['Stories']
    for i, row in enumerate(ws.iter_rows(min_row=3), 0):
        row[1].value, row[2].value = 4.5 * i, 4.5 * (i + 1)
    for row in wb['Sections'].iter_rows(min_row=3):
        row[3].value, row[4].value, row[5].value = 'rect', 450, 2000
    # combos sheet from the headerless ETABS combo paste (as a designer would copy it)
    hc = AIO._headerless_combos(os.path.join(tdir, 'load combination.xlsx'))
    cases = sorted({r['Load Name'] for r in hc}); names = sorted({r['Name'] for r in hc})
    del wb['Combos']; wc = wb.create_sheet('Combos'); wc.append(['note']); wc.append(['combo'] + cases)
    for nme in names:
        f = {r['Load Name']: r['SF'] for r in hc if r['Name'] == nme}
        wc.append([nme] + [f.get(c) for c in cases])
    wb.save(tp)
    tpl = INP.read_template(tp)
    p1 = AIO.prepare({'seismic': (t, u)}, tpl)
    L_ok = all((l, s) in (p1['scS']['L_keys']) for l, s in p1['label_story'])
    check('C1 workbook gives combos, sizes, Lu, story order', not p1['undefined'] and p1['scS']['sec_labels'] == set(p1['labels']) and L_ok
          and p1['stories'][0] == '1st' and p1['stories'][-1] == '20th', str(p1['undefined'][:3]))
    dC = os.path.join(tmp, 'C'); os.makedirs(dC)
    cfgC = AIO.build_cfg(p1, dC)
    resC, outC = AIO.run_engine(dC, cfgC)
    check("C2 f'c fixed 560 below 4th, 400 from 4th (Criteria)", resC['fc_of']['3rd'] == 560 and resC['fc_of']['4th'] == 400)
    check('C2 k 0.9 and βd 0.888 from Criteria', all(abs(q['k'] - 0.9) < 1e-9 and abs(q['betad'] - 0.888) < 1e-9 for r in resC['results'] for q in r['per']))
    check('C2 Lu = story height 4.5 m from Stories', all(abs(q['Lu'] - 450) < 1e-6 for r in resC['results'] for q in r['per']))
    check('C2 combo check flags U_A_00_02 (factors 1.2/1.2/1.6, name says 1.4/1.7)',
          any('U_A_00_02' in n and 'do not match' in n for n in resC['notes']))
    main_f, per = AIO.files(outC)
    names_ = [n for _, n, _ in main_f]
    check('C3 all project outputs written', all(any(n.endswith(x) for n in names_) for x in
          ('_Decision.xlsx', '_Calc.pdf', '_Backup.xlsx', '_DRAFT.dxf', '_DRAFT.pdf', '_Schedule_DRAFT.xlsx')), str(names_))
    marks = set(resC['marks'].values())
    check('C3 per-mark calc PDF + back-up Excel for every mark', set(per) == marks and all(len(v) == 2 for v in per.values()), f'{len(per)} / {len(marks)}')
    png = AIO.dxf_png([p for p, n, _ in main_f if n.endswith('.dxf')][0], dpi=40)
    check('C3 DXF renders', len(png) > 5000, f'{len(png)} bytes')
    r0 = next(r for r in resC['results'] if r.get('ref'))
    f3, f2, fc, gov = AIO.surface_figures(AIO.mark_band_eval(outC, r0['mark'], tuple(r0['band'])))
    dcs = [q['dc'] for r in resC['results'] if r['mark'] == r0['mark'] and tuple(r['band']) == tuple(r0['band']) for q in r['per']]
    check('C3 interaction surface: governing D/C = engine result', abs(gov['ev']['dcmax'] - max(dcs)) < 1e-9, f"{gov['ev']['dcmax']:.4f}")
    # round trip: the prefilled workbook written from the current input reads back the same
    tp2 = os.path.join(tmp, 'again.xlsx'); open(tp2, 'wb').write(INP.make_template(p1['found'], tpl))
    tpl2 = INP.read_template(tp2)
    check('C4 input workbook round trip (download -> upload gives the same config)', INP.to_cfg(tpl2) == INP.to_cfg(tpl))
else:
    print('SKIP C (pass the folder with "Column Load.xlsx" and "load combination.xlsx")')

# ---------------------------------------------------------------- D workbook overrides an ETABS size
t, u, _ = AIO.read_upload(os.path.join(tmp, 'full_red.xlsx'))
ts, us, _ = AIO.read_upload(os.path.join(tmp, 'seis.xlsx'))
tpl = dict(criteria={'project.job': 'OV'}, stories=[], columns=[], sections=[['C2', None, None, 'rect', 700, 700]], combos_rows=[],
           combo_cases=[], lap_rows=[], exclude=['U_B_99'], approve=[], issues=[])
pD = AIO.prepare({'seismic': (ts, us), 'full_reduced': (t, u)}, tpl)
check('D merge report says the workbook replaces ETABS sizes', any('Sections replace' in r for r in pD['merge_rep']))
dD = os.path.join(tmp, 'D'); os.makedirs(dD)
resD, _ = AIO.run_engine(dD, AIO.build_cfg(pD, dD))
c2 = [r for r in resD['results'] if r['label'] == 'C2']
c1 = [r for r in resD['results'] if r['label'] == 'C1']
check('D C2 designed as 700×700 (input workbook), C1 keeps the ETABS 400×800', c2 and c1 and all(r['b'] == 70 and r['h'] == 70 for r in c2)
      and all((r['b'], r['h']) == (40, 80) for r in c1), str({(r['label'], r['b'], r['h']) for r in resD['results']}))
check('D basement still assembled', bool(resD['input_report'].get('basement_stories')))
shutil.rmtree(tmp, ignore_errors=True)
print('ALL PASS' if not fails else f'{len(fails)} FAILED: {fails}')
sys.exit(1 if fails else 0)
